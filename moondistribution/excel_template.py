"""Excel template for corps without Auth access (e.g. behind the Chinese firewall).

Export: one .xlsx per round with all moons of the pool and their details; the corp enters its ticker
and an x next to every moon it wants and sends the file back to a Mining Officer.
Import: the officer uploads the returned files on the round page; each file becomes the application of
the corp named in it. Everything happens behind the Auth login – there is no public page.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.db import transaction
from django.utils.translation import gettext as _

from allianceauth.services.hooks import get_extension_logger

from .models import Application, ApplicationWish, DistributionSettings, Moon
from .probation import probation_status
from .providers import alliance_corporations

logger = get_extension_logger(__name__)

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB per uploaded file
META_SHEET = "meta"
TICKER_CELL = "B3"
HEADER_ROW = 6
FIRST_MOON_ROW = 7
ID_COLUMN = 8  # column H
MARK_EMPTY = {"", "0", "-", "no", "nein", "否", "n"}

HEADERS = [
    "Wanted (x) / 想要 (x)", "Region / 星域", "Moon / 卫星", "Rarity / 稀有度", "Best % / 最佳 %",
    "Composition / 成分", "Current owner / 当前所有者", "ID",
]
RARITY_FILL = {"R64": "F28B82", "R32": "FFA94D", "R16": "69DB7C", "R8": "74A9F7", "R4": "FFFFFF"}


def _pool(round_obj):
    """Moons of the round that can be applied for (R64 excluded if set), best first per region."""
    qs = round_obj.moons.select_related("owner_corporation")
    if DistributionSettings.get_solo().exclude_r64:
        qs = qs.exclude(rarity=Moon.Rarity.R64)
    moons = list(qs)
    moons.sort(key=lambda m: (m.region_name or "~", -m.rarity_sort, m.name))
    return moons


def export_round(round_obj) -> bytes:
    """Build the Excel template of a round: header with ticker cell and instructions, one row per moon."""
    settings = DistributionSettings.get_solo()
    wb = Workbook()
    ws = wb.active
    ws.title = "Moons 卫星"
    bold = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFE066")

    ws["A1"] = f"Moon Distribution – {round_obj.name}"
    ws["A1"].font = Font(bold=True, size=14)
    end = round_obj.application_end.strftime("%Y-%m-%d %H:%M") if round_obj.application_end else "–"
    ws["A2"] = f"Applications until / 申请截止: {end} EVE"
    ws["A3"] = "Corp ticker / 军团代号:"
    ws["A3"].font = bold
    ws[TICKER_CELL].fill = yellow
    ws["A4"] = (
        f"Enter your corp ticker in the yellow cell and put an x in column A for every moon your corp wants "
        f"(max. {settings.max_wishes}). Send the file back to your Mining Officer. / "
        f"在黄色单元格中填写军团代号，并在 A 列为你的军团想要的每个卫星填写 x（最多 {settings.max_wishes} 个），"
        f"然后把文件发回给你的 Mining Officer。"
    )
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:H4")
    ws.row_dimensions[4].height = 48

    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="D9D9D9")

    for row, moon in enumerate(_pool(round_obj), start=FIRST_MOON_ROW):
        ws.cell(row=row, column=1).fill = yellow
        ws.cell(row=row, column=2, value=moon.region_name or "")
        ws.cell(row=row, column=3, value=moon.name)
        rarity = ws.cell(row=row, column=4, value=moon.rarity or "?")
        if moon.rarity in RARITY_FILL:
            rarity.fill = PatternFill("solid", fgColor=RARITY_FILL[moon.rarity])
        ws.cell(row=row, column=5, value=moon.best_share)
        ws.cell(row=row, column=6, value=", ".join(
            f"{c['ore']} {c['percent']:g}%" if c.get("percent") is not None else c["ore"]
            for c in moon.composition_display
        ))
        ws.cell(row=row, column=7, value=moon.owner_display or "")
        ws.cell(row=row, column=ID_COLUMN, value=moon.pk).font = Font(color="999999")

    for col, width in enumerate((16, 14, 30, 12, 12, 60, 28, 8), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=FIRST_MOON_ROW, column=1)
    ws.auto_filter.ref = f"A{HEADER_ROW}:H{max(HEADER_ROW, ws.max_row)}"

    meta = wb.create_sheet(META_SHEET)
    meta["A1"], meta["B1"] = "app", "aa-moondistribution"
    meta["A2"], meta["B2"] = "round", round_obj.pk
    meta.sheet_state = "hidden"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _ticker(value) -> str:
    """'[ABC] ', 'abc' → 'ABC'."""
    return str(value or "").strip().strip("[]").strip().upper()


def read_template(uploaded, round_obj) -> dict:
    """Read one returned template. Returns {"ticker", "moon_ids"} or {"error"} – nothing is saved here."""
    if uploaded.size > MAX_FILE_SIZE:
        return {"error": _("file larger than 5 MB")}
    try:
        wb = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception:  # not an xlsx, damaged file …
        return {"error": _("not an Excel file (.xlsx)")}
    if META_SHEET not in wb.sheetnames:
        return {"error": _("not an Excel template of this app")}
    meta = {row[0]: row[1] for row in wb[META_SHEET].iter_rows(min_row=1, max_row=5, values_only=True) if row}
    if str(meta.get("round")) != str(round_obj.pk):
        return {"error": _("template of another round – please use the template of this round")}
    ws = wb.worksheets[0]
    ticker = _ticker(ws[TICKER_CELL].value)
    moon_ids = []
    for row in ws.iter_rows(min_row=FIRST_MOON_ROW, values_only=True):
        if not row or len(row) < ID_COLUMN:
            continue
        mark = str(row[0] if row[0] is not None else "").strip().lower()
        if mark in MARK_EMPTY:
            continue
        try:
            moon_ids.append(int(row[ID_COLUMN - 1]))
        except (TypeError, ValueError):
            continue
    wb.close()
    return {"ticker": ticker, "moon_ids": moon_ids}


def import_templates(round_obj, files, user) -> list:
    """Import the returned templates of a round. Returns one result per file:
    {"file", "ok": bool, "corp": name or "", "count": n, "message": str}.

    Same rules as a normal application: alliance corp, pool moons only (no R64), max. wishes, probation.
    The application is marked as entered by an officer (rated without the member criteria if the corp
    has no player in the Auth).
    """
    settings = DistributionSettings.get_solo()
    pool = {m.pk: m for m in _pool(round_obj)}
    corps = {c.corporation_ticker.upper(): c for c in alliance_corporations()}
    results = []
    for uploaded in files:
        res = {"file": uploaded.name, "ok": False, "corp": "", "count": 0, "message": ""}
        data = read_template(uploaded, round_obj)
        if "error" in data:
            res["message"] = data["error"]
            results.append(res)
            continue
        corp = corps.get(data["ticker"])
        if corp is None:
            res["message"] = _("ticker '%(ticker)s' is not a corp of the alliance") % {"ticker": data["ticker"] or "–"}
            results.append(res)
            continue
        res["corp"] = corp.corporation_name
        status = probation_status(corp, months=settings.probation_months)
        if status["blocked"]:
            res["message"] = _("still in its probation until %(date)s") % {"date": status["ends_at"].strftime("%Y-%m-%d")}
            results.append(res)
            continue
        chosen = [pool[mid] for mid in dict.fromkeys(data["moon_ids"]) if mid in pool]
        if not chosen:
            res["message"] = _("no moon marked with x")
            results.append(res)
            continue
        if len(chosen) > settings.max_wishes:
            res["message"] = _("%(n)s moons marked – at most %(max)s allowed") % {"n": len(chosen), "max": settings.max_wishes}
            results.append(res)
            continue
        with transaction.atomic():
            application, _created = Application.objects.select_for_update().get_or_create(
                round=round_obj, corporation=corp, defaults={"submitted_by": user},
            )
            application.submitted_by = user
            application.entered_by_officer = True
            if not application.comment:
                application.comment = f"Imported from the Excel template ({uploaded.name})"[:2000]
            application.save()
            application.wishes.all().delete()
            ApplicationWish.objects.bulk_create([
                ApplicationWish(application=application, moon=moon, priority=1) for moon in chosen
            ])
        ignored = len(data["moon_ids"]) - len(chosen)
        res.update(ok=True, count=len(chosen), message=_("%(n)s marked rows ignored (not in the pool)") % {"n": ignored} if ignored else "")
        logger.info("Template import: corp %s, %s moons, round %s, by user %s",
                    corp.corporation_id, len(chosen), round_obj.pk, user.pk)
        results.append(res)
    return results
